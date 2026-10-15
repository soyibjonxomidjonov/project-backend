import uuid

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import async_session, get_db
from app.deps import get_current_user
from app.models.document import Document, DocumentChunk
from app.models.user import User
from app.schemas.knowledge import (
    ChunkResult,
    DocumentResponse,
    KnowledgeSearchRequest,
    KnowledgeTextRequest,
)
from app.services.embeddings import embed_texts
from app.services.rag import chunk_text, search_similar

router = APIRouter(prefix="/knowledge", tags=["knowledge"])


async def _process_document(document_id: uuid.UUID, text: str):
    chunks = chunk_text(text)
    if not chunks:
        async with async_session() as db:
            doc = (await db.execute(select(Document).where(Document.id == document_id))).scalar_one()
            doc.status = "failed"
            await db.commit()
        return

    embeddings = await embed_texts(chunks)

    async with async_session() as db:
        doc = (await db.execute(select(Document).where(Document.id == document_id))).scalar_one()
        for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
            db.add(
                DocumentChunk(
                    document_id=document_id,
                    content=chunk,
                    chunk_index=i,
                    embedding=embedding,
                )
            )
        doc.chunk_count = len(chunks)
        doc.status = "ready"
        await db.commit()


IMAGE_MIME_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif"}


def _is_image(mime_type: str) -> bool:
    return mime_type in IMAGE_MIME_TYPES


def _extract_text(file_bytes: bytes, mime_type: str) -> str:
    if mime_type == "application/pdf":
        from io import BytesIO

        from pypdf import PdfReader

        reader = PdfReader(BytesIO(file_bytes))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        from io import BytesIO

        from docx import Document as DocxDoc

        doc = DocxDoc(BytesIO(file_bytes))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    return file_bytes.decode("utf-8", errors="replace")


@router.post("/upload", summary="Upload document for RAG", response_model=DocumentResponse)
async def upload_document(
    file: UploadFile,
    background_tasks: BackgroundTasks,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    content = await file.read()
    mime = file.content_type or "text/plain"

    if _is_image(mime):
        from app.services.gemini import analyze_image

        try:
            text = await analyze_image(content, mime)
        except Exception:
            raise HTTPException(400, "Rasmni tahlil qilishda xatolik")
    else:
        try:
            text = _extract_text(content, mime)
        except Exception:
            raise HTTPException(400, "Faylni o'qishda xatolik")

    doc = Document(
        title=file.filename or "Untitled",
        filename=file.filename,
        mime_type=mime,
        owner_id=user.id,
        status="processing",
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    background_tasks.add_task(_process_document, doc.id, text)
    return DocumentResponse.model_validate(doc)


@router.post("/text", summary="Add text directly to knowledge base", response_model=DocumentResponse)
async def add_text(
    body: KnowledgeTextRequest,
    background_tasks: BackgroundTasks,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    doc = Document(
        title=body.title,
        mime_type="text/plain",
        owner_id=user.id,
        status="processing",
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    background_tasks.add_task(_process_document, doc.id, body.content)
    return DocumentResponse.model_validate(doc)


@router.get("/documents", summary="List knowledge documents")
async def list_documents(
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(Document).order_by(Document.created_at.desc()))
    docs = result.scalars().all()
    return {"items": [DocumentResponse.model_validate(d) for d in docs]}


@router.delete("/documents/{document_id}", summary="Delete document", status_code=204)
async def delete_document(
    document_id: uuid.UUID,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(Document).where(Document.id == document_id))
    doc = result.scalar_one_or_none()
    if doc is None:
        raise HTTPException(404, "Document not found")
    if doc.owner_id != user.id and user.role != "admin":
        raise HTTPException(403, "Not allowed")
    await db.delete(doc)
    await db.commit()


@router.post("/search", summary="Search knowledge base", response_model=list[ChunkResult])
async def search_knowledge(
    body: KnowledgeSearchRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    results = await search_similar(db, body.query)
    return [ChunkResult(**r) for r in results]
